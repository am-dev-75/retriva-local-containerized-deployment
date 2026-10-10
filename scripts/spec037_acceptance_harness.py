# Copyright (C) 2026 Retriva. All rights reserved.
# Proprietary — Retriva Pro extension. See LICENSE.retriva-pro.
#
# Spec 037 Phase 5 — isolated deployed acceptance harness.
#
# Runs INSIDE the deployed retriva-core (pro) container against the real
# isolated PostgreSQL, the real deployed settings, the real stores/stage,
# and the real Creditsafe client/auth/resolver/scoring code.  HTTP calls
# leave the client through an injected ``httpx.MockTransport`` (the same
# provider-neutral seam the test-suite uses); no production escape hatch
# exists and the production HTTPS/base-URL validation is untouched.
#
# All data is synthetic.  No real credentials, companies, or reports.

from __future__ import annotations

import argparse
import json
import sys
import uuid
from datetime import datetime, timedelta, timezone
from decimal import Decimal
from typing import Any, Dict, List, Optional

import httpx

TENANT_A = "cust_s037_a"
TENANT_B = "cust_s037_b"

A1 = "org_s037_a1"
A2 = "org_s037_a2"
A3 = "org_s037_a3"
A4 = "org_s037_a4"
B1 = "org_s037_b1"

UTC = timezone.utc


def out(payload: Any) -> None:
    print(json.dumps(payload, indent=2, default=str))


def pool():
    from retriva_crm_assistant.postgres.config import (
        get_postgres_settings,
    )
    from retriva_crm_assistant.postgres.pool import get_pool
    settings = get_postgres_settings()
    return get_pool("application"), settings


def settings_summary() -> Dict[str, Any]:
    from retriva_crm_assistant.config import crm_settings
    return {
        "creditsafe_enabled": bool(crm_settings.creditsafe_enabled),
        "creditsafe_environment": crm_settings.creditsafe_environment,
        "budget_company_search": crm_settings.creditsafe_budget_company_search,
        "budget_credit_report": crm_settings.creditsafe_budget_credit_report,
        "freshness_days": crm_settings.creditsafe_freshness_days,
        "allow_name_only_search": bool(
            crm_settings.creditsafe_allow_name_only_search),
        "has_username": bool(crm_settings.creditsafe_username),
        "has_password": bool(
            crm_settings.creditsafe_password.get_secret_value()
            if hasattr(crm_settings.creditsafe_password, "get_secret_value")
            else crm_settings.creditsafe_password),
    }


# ---------------------------------------------------------------------------
# Synthetic mock transport (constitution-compliant: injected at the
# provider-neutral transport seam; records every request for evidence).
# ---------------------------------------------------------------------------

def _jwt() -> str:
    return "synthetic.jwt.signature"


class MockCreditsafe:
    def __init__(self) -> None:
        self.routes: Dict[str, List[Any]] = {}
        self.calls: List[Dict[str, Any]] = []

    def route(self, method: str, path: str, *responders) -> "MockCreditsafe":
        self.routes.setdefault(f"{method.upper()} {path}",
                               []).extend(responders)
        return self

    def json_route(self, method: str, path: str, payload: Any,
                   status: int = 200) -> "MockCreditsafe":
        def _r(_req: httpx.Request, body: Any = payload,
               code: int = status) -> httpx.Response:
            return httpx.Response(code, json=body)
        return self.route(method, path, _r)

    def sequence(self, method: str, path: str, *steps: Any
                 ) -> "MockCreditsafe":
        for step in steps:
            if isinstance(step, int):
                def _s(_req: httpx.Request, code: int = step
                       ) -> httpx.Response:
                    return httpx.Response(code, json={"error": "synthetic"})
                self.route(method, path, _s)
            else:
                self.json_route(method, path, step)
        return self

    def transport(self) -> httpx.MockTransport:
        def _handler(request: httpx.Request) -> httpx.Response:
            path = request.url.path
            if path.startswith("/v1/"):
                path = path[len("/v1"):]
            key = f"{request.method.upper()} {path}"
            self.calls.append({
                "method": request.method.upper(),
                "path": path,
                "params": dict(request.url.params),
            })
            queue = self.routes.get(key)
            if not queue:
                return httpx.Response(404, json={"error": "unrouted"})
            responder = queue.pop(0) if len(queue) > 1 else queue[0]
            return responder(request)
        return httpx.MockTransport(_handler)

    def summary(self) -> Dict[str, Any]:
        counts: Dict[str, int] = {}
        for call in self.calls:
            counts[f"{call['method']} {call['path']}"] = counts.get(
                f"{call['method']} {call['path']}", 0) + 1
        return {"total": len(self.calls), "by_route": counts,
                "calls": [f"{c['method']} {c['path']}" for c in self.calls]}


def company_row(*, connect_id: str, reg_no: str, name: str,
                country: str = "IT", postcode: str = "00100",
                city: str = "Roma", street: str = "Via Sintetica 1"
                ) -> Dict[str, Any]:
    # Address mirrors the synthetic fixture address (Via Sintetica 1,
    # Roma, 00100) so resolver comparisons match exactly; organizations
    # without an address row skip address comparison entirely.
    return {
        "id": connect_id, "country": country, "regNo": reg_no,
        "name": name,
        "address": {"simpleValue": street, "street": street,
                    "city": city, "postCode": postcode},
    }


def search_payload(companies: List[Dict[str, Any]]) -> Dict[str, Any]:
    return {"correlationId": "corr-search", "messages": [],
            "totalSize": len(companies), "companies": companies}


def report_payload(*, connect_id: str, value: Any = 80,
                   min_value: Any = 0, max_value: Any = 100,
                   failed_sections: Optional[List[str]] = None
                   ) -> Dict[str, Any]:
    report: Dict[str, Any] = {
        "companyId": connect_id,
        "companySummary": {"name": "Synthetic"},
    }
    if failed_sections is None:
        report["creditScore"] = {
            "currentCreditRating": {
                "commonValue": "B",
                "providerValue": {"value": value,
                                  "minValue": min_value,
                                  "maxValue": max_value},
                "pod": 1.25,
            },
            "currentCreditLimit": {"currency": "EUR", "value": 10000},
        }
    return {
        "correlationId": "corr-report",
        "failedSections": failed_sections or [],
        "report": report,
        "companyId": connect_id,
        "dateOfOrder": "2026-10-01T00:00:00Z",
        "language": "en",
        "userId": "synthetic",
    }


def success_mock(*, connect_id: str, reg_no: str, name: str,
                 value: Any = 80) -> MockCreditsafe:
    mock = MockCreditsafe()
    mock.json_route("POST", "/authenticate", {"token": _jwt()})
    mock.json_route("GET", "/companies", search_payload([
        company_row(connect_id=connect_id, reg_no=reg_no, name=name)]))
    mock.json_route("GET", f"/companies/{connect_id}",
                    report_payload(connect_id=connect_id, value=value))
    return mock


def make_mock_provider(mock: MockCreditsafe):
    from retriva_crm_assistant.config import crm_settings
    from retriva_crm_assistant.creditsafe.budget import (
        CreditsafeBudgetGuard,
    )
    from retriva_crm_assistant.creditsafe.client import (
        CreditsafeHttpClient,
    )
    from retriva_crm_assistant.creditsafe.config import CreditsafeConfig
    from retriva_crm_assistant.creditsafe.provider import (
        CreditsafeFinancialRiskProvider,
    )
    from retriva_crm_assistant.intelligence_metrics import (
        IntelligenceMetrics,
    )
    config = CreditsafeConfig.from_settings(crm_settings)
    metrics = IntelligenceMetrics()
    guard = CreditsafeBudgetGuard(
        company_search_limit=config.budget_company_search,
        credit_report_limit=config.budget_credit_report,
        metrics=metrics)
    client = CreditsafeHttpClient(
        config, transport=mock.transport(), budget=guard, metrics=metrics)
    provider = CreditsafeFinancialRiskProvider(
        config, client=client, metrics=metrics)
    return provider, guard


def obs_to_dict(o: Any) -> Optional[Dict[str, Any]]:
    if o is None:
        return None
    return {
        "observation_id": o.observation_id,
        "score_status": getattr(o.score_status, "value", o.score_status),
        "normalized_score": str(o.normalized_score)
        if o.normalized_score is not None else None,
        "provider_scale_value": str(o.provider_scale_value)
        if o.provider_scale_value is not None else None,
        "provider_scale_min": str(o.provider_scale_min)
        if o.provider_scale_min is not None else None,
        "provider_scale_max": str(o.provider_scale_max)
        if o.provider_scale_max is not None else None,
        "connect_id": o.connect_id,
        "observed_at": o.observed_at,
        "retrieved_at": o.retrieved_at,
        "accepted_as_current": o.accepted_as_current,
        "failure_category": o.failure_category,
        "identity_verification": o.identity_verification,
    }


def state_to_dict(s: Any) -> Optional[Dict[str, Any]]:
    if s is None:
        return None
    return {
        "accepted_observation_id": s.accepted_observation_id,
        "last_status": s.last_status,
        "last_failure_category": s.last_failure_category,
        "last_attempt_at": s.last_attempt_at,
        "last_success_at": s.last_success_at,
        "connect_id": s.connect_id,
    }


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

def _insert_org(p, tenant: str, org: str, name: str) -> None:
    with p.transaction(tenant_id=tenant) as conn:
        with conn.cursor() as cur:
            cur.execute(
                "INSERT INTO business.organizations ("
                "organization_id, tenant_id, canonical_legal_name, "
                "normalized_legal_name, country_code) "
                "VALUES (%s, %s, %s, %s, 'IT')",
                (org, tenant, name, name.lower()))


def _insert_address(p, tenant: str, org: str) -> None:
    with p.transaction(tenant_id=tenant) as conn:
        with conn.cursor() as cur:
            cur.execute(
                "INSERT INTO business.organization_addresses ("
                "organization_address_id, tenant_id, organization_id, "
                "address_type, address_line_1, postal_code, city, "
                "normalized_address_hash) "
                "VALUES (%s, %s, %s, 'REGISTERED', 'Via Sintetica 1', "
                "'00100', 'Roma', %s)",
                (f"oa_{uuid.uuid4().hex}", tenant, org,
                 uuid.uuid4().hex))


def _insert_role(p, tenant: str, org: str, role: str = "CUSTOMER") -> None:
    with p.transaction(tenant_id=tenant) as conn:
        with conn.cursor() as cur:
            cur.execute(
                "INSERT INTO business.organization_roles ("
                "organization_role_id, tenant_id, organization_id, "
                "role_type, role_status) "
                "VALUES (%s, %s, %s, %s, 'ACTIVE')",
                (f"or_{uuid.uuid4().hex}", tenant, org, role))


def _insert_identifier(p, tenant: str, org: str, kind: str,
                       value: str) -> None:
    with p.transaction(tenant_id=tenant) as conn:
        with conn.cursor() as cur:
            cur.execute(
                "INSERT INTO business.organization_identifiers ("
                "organization_identifier_id, tenant_id, organization_id, "
                "identifier_type, normalized_identifier_value) "
                "VALUES (%s, %s, %s, %s, %s)",
                (f"oi_{uuid.uuid4().hex}", tenant, org, kind, value))


FIXTURES = {
    (TENANT_A, A1): ("Synthetic Alfa Srl", "RN-S037-A1",
                     "IT-S037-A1-VAT", True),
    (TENANT_A, A2): ("Synthetic Beta Srl", "RN-S037-A2",
                     "IT-S037-A2-VAT", True),
    (TENANT_A, A3): ("Synthetic Gamma Srl", "RN-S037-A3",
                     "IT-S037-A3-VAT", False),
    (TENANT_A, A4): ("Synthetic Delta Srl", "RN-S037-A4",
                     "IT-S037-A4-VAT", False),
    (TENANT_B, B1): ("Synthetic Epsilon Srl", "RN-S037-B1",
                     "IT-S037-B1-VAT", True),
}


def cmd_seed_fixtures(_args) -> int:
    p, _ = pool()
    for (tenant, org), (name, rn, vat, with_address) in FIXTURES.items():
        _insert_org(p, tenant, org, name)
        if with_address:
            _insert_address(p, tenant, org)
        _insert_identifier(p, tenant, org, "REGISTRATION_NUMBER", rn)
        _insert_identifier(p, tenant, org, "VAT_ID", vat)
        if tenant == TENANT_A and org in (A1, A2):
            _insert_role(p, tenant, org, "CUSTOMER")
    counts = {}
    for tenant in (TENANT_A, TENANT_B):
        with p.transaction(tenant_id=tenant) as conn:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT count(*) AS n FROM business.organizations "
                    "WHERE tenant_id = %s", (tenant,))
                counts[tenant] = cur.fetchone()["n"]
    out({"tenants": [TENANT_A, TENANT_B],
         "organizations": {t: sorted(o for (tt, o) in FIXTURES
                                     if t == tt) for t in
                           (TENANT_A, TENANT_B)},
         "org_counts": counts})
    return 0


# ---------------------------------------------------------------------------
# Stage scenarios
# ---------------------------------------------------------------------------

def _evidence_stage(p, tenant: str, org: str, mock: MockCreditsafe,
                    guard: Any, outcome: Any) -> Dict[str, Any]:
    from retriva_crm_assistant.postgres.qualification import (
        financial_scores as fstore,
    )
    return {
        "outcome": outcome.to_result(),
        "mock_calls": mock.summary(),
        "budget_snapshot": guard.snapshot() if guard else None,
        "state": state_to_dict(fstore.get_state(p, tenant, org)),
        "current": obs_to_dict(
            fstore.get_current_observation(p, tenant, org)),
        "history": [obs_to_dict(o) for o in
                    fstore.list_observations(p, tenant, org)],
    }


def cmd_scenario_success(_args) -> int:
    from retriva_crm_assistant.postgres.qualification.enrichment import (
        creditsafe_stage as stage,
    )
    from retriva_crm_assistant.postgres.qualification import (
        financial_scores as fstore,
    )
    p, _ = pool()
    mock = success_mock(connect_id="CS-S037-A1", reg_no="RN-S037-A1",
                        name="Synthetic Alfa Srl", value=80)
    provider, guard = make_mock_provider(mock)

    r1 = stage.run_creditsafe_stage(p, TENANT_A, A1, provider=provider)
    run1 = _evidence_stage(p, TENANT_A, A1, mock, guard, r1)

    r2 = stage.run_creditsafe_stage(p, TENANT_A, A1, provider=provider)
    run2 = {"outcome": r2.to_result(), "mock_calls": mock.summary(),
            "budget_snapshot": guard.snapshot()}

    r3 = stage.run_creditsafe_stage(p, TENANT_A, A1, provider=provider,
                                    forced_refresh=True)
    run3 = _evidence_stage(p, TENANT_A, A1, mock, guard, r3)

    out({
        "scenario": "success-reuse-force",
        "settings": settings_summary(),
        "run1": run1,
        "run2": run2,
        "run3": run3,
        "checks": {
            "run1_refreshed": r1.status == stage.STATUS_REFRESHED,
            "run2_reused": r2.status == stage.STATUS_REUSED,
            "run2_no_new_calls": (
                run2["mock_calls"]["total"]
                == run1["mock_calls"]["total"]),
            "run2_no_budget_consumed": (
                run2["budget_snapshot"] == run1["budget_snapshot"]),
            "run3_refreshed": r3.status == stage.STATUS_REFRESHED,
            "run3_new_observation": (
                r3.observation_id != r1.observation_id),
            "history_retains_run1": any(
                o.observation_id == r1.observation_id
                for o in fstore.list_observations(p, TENANT_A, A1)),
            "exact_decimal_score": str(
                fstore.get_current_observation(
                    p, TENANT_A, A1).normalized_score) == "0.8",
            "provider_scale_kept": (
                str(run3["current"]["provider_scale_value"]) == "80"),
            "connect_id_attached_once": _count_connect_ids(
                p, TENANT_A, A1) == 1,
        },
    })
    return 0


def _count_connect_ids(p, tenant: str, org: str) -> int:
    with p.transaction(tenant_id=tenant) as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT count(*) AS n FROM business.organization_identifiers "
                "WHERE tenant_id = %s AND organization_id = %s AND "
                "identifier_type = 'CREDITSAFE_CONNECT_ID'",
                (tenant, org))
            return cur.fetchone()["n"]


def cmd_scenario_zero_budget(_args) -> int:
    from retriva_crm_assistant.postgres.qualification.enrichment import (
        creditsafe_stage as stage,
    )
    from retriva_crm_assistant.postgres.qualification import (
        financial_scores as fstore,
    )
    p, _ = pool()
    mock = MockCreditsafe()  # any call at all is a failure
    provider, guard = make_mock_provider(mock)
    before = len(fstore.list_observations(p, TENANT_A, A2))
    r = stage.run_creditsafe_stage(p, TENANT_A, A2, provider=provider)
    after = len(fstore.list_observations(p, TENANT_A, A2))
    out({
        "scenario": "zero-budget",
        "settings": settings_summary(),
        "readiness": provider.readiness(),
        "stage_outcome": r.to_result(),
        "mock_calls": mock.summary(),
        "observation_rows_before": before,
        "observation_rows_after": after,
        "state": state_to_dict(fstore.get_state(p, TENANT_A, A2)),
        "checks": {
            "skipped_budget_exhausted":
                r.status == stage.STATUS_SKIPPED_BUDGET_EXHAUSTED,
            "zero_provider_calls": len(mock.calls) == 0,
            "no_rows_created": before == after,
            "no_fabricated_score": fstore.get_current_observation(
                p, TENANT_A, A2) is None,
        },
    })
    return 0


def cmd_scenario_budget_boundary(_args) -> int:
    from retriva_crm_assistant.postgres.qualification.enrichment import (
        creditsafe_stage as stage,
    )
    p, _ = pool()
    mock = success_mock(connect_id="CS-S037-A2", reg_no="RN-S037-A2",
                        name="Synthetic Beta Srl", value=74)
    provider, guard = make_mock_provider(mock)
    r1 = stage.run_creditsafe_stage(p, TENANT_A, A2, provider=provider)
    calls_after_first = mock.summary()
    r2 = stage.run_creditsafe_stage(p, TENANT_A, A2, provider=provider,
                                    forced_refresh=True)
    out({
        "scenario": "budget-boundary-1-1",
        "settings": settings_summary(),
        "run1": r1.to_result(),
        "run1_calls": calls_after_first,
        "run1_budget_snapshot": guard.snapshot(),
        "run2_forced": r2.to_result(),
        "run2_calls": mock.summary(),
        "run2_budget_snapshot": guard.snapshot(),
        "checks": {
            "first_refreshed": r1.status == stage.STATUS_REFRESHED,
            "second_skipped_budget":
                r2.status == stage.STATUS_SKIPPED_BUDGET_EXHAUSTED,
            "no_calls_after_exhaustion": (
                mock.summary()["total"]
                == calls_after_first["total"]),
            "search_consumed_once": (
                guard.snapshot()["company_search"]["limit"] == 1
                and guard.snapshot()["company_search"]["remaining"] == 0),
            "report_consumed_once": (
                guard.snapshot()["credit_report"]["limit"] == 1
                and guard.snapshot()["credit_report"]["remaining"] == 0),
        },
    })
    return 0


def cmd_scenario_failure_recover(_args) -> int:
    from retriva_crm_assistant.postgres.qualification.enrichment import (
        creditsafe_stage as stage,
    )
    from retriva_crm_assistant.postgres.qualification import (
        financial_scores as fstore,
    )
    p, _ = pool()
    before_state = state_to_dict(fstore.get_state(p, TENANT_A, A1))
    before_current = obs_to_dict(
        fstore.get_current_observation(p, TENANT_A, A1))

    # 1. Failure: the credit-score section of the report failed.
    fail_mock = MockCreditsafe()
    fail_mock.json_route("POST", "/authenticate", {"token": _jwt()})
    fail_mock.json_route("GET", "/companies", search_payload([
        company_row(connect_id="CS-S037-A1", reg_no="RN-S037-A1",
                    name="Synthetic Alfa Srl")]))
    fail_mock.json_route("GET", "/companies/CS-S037-A1",
                         report_payload(connect_id="CS-S037-A1",
                                        failed_sections=["CreditScore"]))
    provider, guard = make_mock_provider(fail_mock)
    rf = stage.run_creditsafe_stage(p, TENANT_A, A1, provider=provider,
                                    forced_refresh=True)
    mid_state = state_to_dict(fstore.get_state(p, TENANT_A, A1))
    mid_current = obs_to_dict(
        fstore.get_current_observation(p, TENANT_A, A1))

    # 2. Transient failure: 503 then success (retry within one logical
    #    call; no additional budget unit).
    retry_mock = MockCreditsafe()
    retry_mock.json_route("POST", "/authenticate", {"token": _jwt()})
    retry_mock.json_route("GET", "/companies", search_payload([
        company_row(connect_id="CS-S037-A1", reg_no="RN-S037-A1",
                    name="Synthetic Alfa Srl")]))
    retry_mock.sequence(
        "GET", "/companies/CS-S037-A1", 503,
        report_payload(connect_id="CS-S037-A1", value=75))
    provider2, guard2 = make_mock_provider(retry_mock)
    rr = stage.run_creditsafe_stage(p, TENANT_A, A1, provider=provider2,
                                    forced_refresh=True)

    out({
        "scenario": "failure-preserve-recover",
        "before": {"state": before_state, "current": before_current},
        "failure_run": {
            "outcome": rf.to_result(),
            "calls": fail_mock.summary(),
            "state_after": mid_state,
            "current_after": mid_current,
            "history": [obs_to_dict(o) for o in
                        fstore.list_observations(p, TENANT_A, A1)],
        },
        "retry_run": {
            "outcome": rr.to_result(),
            "calls": retry_mock.summary(),
            "budget_snapshot": guard2.snapshot(),
            "state_after": state_to_dict(
                fstore.get_state(p, TENANT_A, A1)),
            "current_after": obs_to_dict(
                fstore.get_current_observation(p, TENANT_A, A1)),
            "history": [obs_to_dict(o) for o in
                        fstore.list_observations(p, TENANT_A, A1)],
        },
        "checks": {
            "failure_unavailable": rf.status == stage.STATUS_UNAVAILABLE,
            "accepted_link_preserved": (
                mid_state["accepted_observation_id"]
                == before_state["accepted_observation_id"]),
            "last_success_preserved": (
                mid_state["last_success_at"] == before_state["last_success_at"]),
            "no_zero_score_written": all(
                (o.normalized_score is None)
                or (o.normalized_score != Decimal(0))
                for o in fstore.list_observations(p, TENANT_A, A1)),
            "recovery_refreshed": rr.status == stage.STATUS_REFRESHED,
            "retry_within_one_logical_call": (
                retry_mock.summary()["by_route"].get(
                    "GET /companies/CS-S037-A1") == 2
                and guard2.snapshot()["credit_report"]["remaining"] == 3),
        },
    })
    return 0


def cmd_scenario_conflicts(_args) -> int:
    from retriva_crm_assistant.postgres.qualification.enrichment import (
        creditsafe_stage as stage,
    )
    from retriva_crm_assistant.postgres.qualification import (
        financial_scores as fstore,
    )
    p, _ = pool()
    results: Dict[str, Any] = {"scenario": "identifier-conflicts"}

    # Tenant A, org A3: provider returns the connect id already held
    # by A2 IN THE SAME TENANT -> IDENTITY_CONFLICT, no promotion.
    mock = success_mock(connect_id="CS-S037-A2", reg_no="RN-S037-A3",
                        name="Synthetic Gamma Srl", value=90)
    provider, _ = make_mock_provider(mock)
    rc = stage.run_creditsafe_stage(p, TENANT_A, A3, provider=provider)
    results["same_tenant_conflict"] = {
        "outcome": rc.to_result(),
        "calls": mock.summary(),
        "state": state_to_dict(fstore.get_state(p, TENANT_A, A3)),
        "current": obs_to_dict(
            fstore.get_current_observation(p, TENANT_A, A3)),
        "connect_ids_a3": _count_connect_ids(p, TENANT_A, A3),
        "a2_current_still": obs_to_dict(
            fstore.get_current_observation(p, TENANT_A, A2)),
    }

    # Cross-tenant: tenant B org B1 uses the SAME connect id value ->
    # allowed (tenant-scoped uniqueness).
    mock_b = success_mock(connect_id="CS-S037-A2", reg_no="RN-S037-B1",
                          name="Synthetic Epsilon Srl", value=65)
    provider_b, _ = make_mock_provider(mock_b)
    rb = stage.run_creditsafe_stage(p, TENANT_B, B1, provider=provider_b)
    results["cross_tenant_reuse"] = {
        "outcome": rb.to_result(),
        "calls": mock_b.summary(),
        "state": state_to_dict(fstore.get_state(p, TENANT_B, B1)),
        "current": obs_to_dict(
            fstore.get_current_observation(p, TENANT_B, B1)),
        "connect_ids_b1": _count_connect_ids(p, TENANT_B, B1),
    }
    results["checks"] = {
        "conflict_detected": rc.status == stage.STATUS_IDENTITY_CONFLICT,
        "conflict_no_connect_id_attached": (
            results["same_tenant_conflict"]["connect_ids_a3"] == 0),
        "conflict_no_score_promoted": (
            results["same_tenant_conflict"]["current"] is None),
        "a2_score_untouched": (
            results["same_tenant_conflict"]["a2_current_still"]
            is not None),
        "cross_tenant_allowed": (
            rb.status == stage.STATUS_REFRESHED
            and results["cross_tenant_reuse"]["connect_ids_b1"] == 1),
    }
    out(results)
    return 0


def cmd_scenario_stale_and_tenant_isolation(_args) -> int:
    from retriva_crm_assistant.postgres.qualification.enrichment import (
        creditsafe_stage as stage,
    )
    from retriva_crm_assistant.postgres.qualification import (
        financial_scores as fstore,
    )
    from retriva_crm_assistant.models import (
        FinancialScoreObservation, FinancialScoreStatus,
    )
    p, _ = pool()

    # A4: synthetic STALE success (120 days old) recorded directly.
    old = (datetime.now(UTC) - timedelta(days=120)).isoformat()
    fstore.record_success(p, TENANT_A, A4, FinancialScoreObservation(
        score_status=FinancialScoreStatus.AVAILABLE,
        normalized_score=Decimal("0.9"),
        observed_at=old))
    stale_before = obs_to_dict(
        fstore.get_current_observation(p, TENANT_A, A4))

    # Without force the stage treats it as stale and refreshes via mock.
    mock = success_mock(connect_id="CS-S037-A4", reg_no="RN-S037-A4",
                        name="Synthetic Delta Srl", value=70)
    provider, _ = make_mock_provider(mock)
    r = stage.run_creditsafe_stage(p, TENANT_A, A4, provider=provider)
    stale_after = obs_to_dict(
        fstore.get_current_observation(p, TENANT_A, A4))

    # Tenant isolation at the store level (application role + RLS):
    # tenant B context must not see tenant A rows and vice versa.
    cross = {}
    for tenant, foreign_org in ((TENANT_B, A1), (TENANT_A, B1)):
        cross[f"{tenant}->{foreign_org}"] = {
            "current": obs_to_dict(fstore.get_current_observation(
                p, tenant, foreign_org)),
            "state": state_to_dict(fstore.get_state(
                p, tenant, foreign_org)),
            "history_len": len(fstore.list_observations(
                p, tenant, foreign_org)),
        }
    out({
        "scenario": "stale-refresh-and-tenant-isolation",
        "stale_before": stale_before,
        "stale_run": {"outcome": r.to_result(), "calls": mock.summary()},
        "stale_after": stale_after,
        "cross_tenant_reads": cross,
        "checks": {
            "stale_was_marked_stale_input": stale_before["observed_at"]
            == old,
            "stale_refreshed": r.status == stage.STATUS_REFRESHED,
            "new_observation_replaces": (
                stale_after["observation_id"]
                != stale_before["observation_id"]),
            "cross_tenant_reads_empty": all(
                v["history_len"] == 0 for v in cross.values()),
        },
    })
    return 0


def cmd_query(args) -> int:
    from retriva_crm_assistant.postgres.qualification import (
        financial_scores as fstore,
    )
    from retriva_crm_assistant.postgres.qualification import (
        financial_settings as fsettings,
    )
    p, _ = pool()
    what = args.what
    if what == "threshold":
        out({"tenant_a": fsettings.get_effective_settings(p, TENANT_A),
             "tenant_b": fsettings.get_effective_settings(p, TENANT_B)})
    elif what == "states":
        payload = {}
        for tenant, orgs in ((TENANT_A, (A1, A2, A3, A4)),
                             (TENANT_B, (B1,))):
            for org in orgs:
                payload[f"{tenant}/{org}"] = {
                    "state": state_to_dict(fstore.get_state(
                        p, tenant, org)),
                    "history": [obs_to_dict(o) for o in
                                fstore.list_observations(p, tenant, org)],
                }
        out(payload)
    elif what == "rls":
        out(_rls_evidence(p))
    elif what == "provider":
        from retriva_crm_assistant.postgres.qualification.enrichment import (
            creditsafe_stage as stage,
        )
        provider = stage.get_creditsafe_provider()
        out({"settings": settings_summary(),
             "readiness": provider.readiness()})
    elif what == "audit":
        with p.transaction(tenant_id=args.tenant or TENANT_A) as conn:
            with conn.cursor() as cur:
                cur.execute(
                    "SELECT action, entity_type, entity_id, actor_id, "
                    "before_summary, after_summary, event_timestamp "
                    "FROM audit.events WHERE action LIKE '%FINANCIAL%' "
                    "OR action LIKE '%IDENTIFIER%' "
                    "ORDER BY event_timestamp DESC LIMIT 25")
                out([dict(r) for r in cur.fetchall()])
    else:
        raise SystemExit(f"unknown query {what!r}")
    return 0


def _rls_evidence(p) -> Dict[str, Any]:
    with p.transaction(tenant_id=TENANT_A) as conn:
        with conn.cursor() as cur:
            cur.execute(
                "SELECT c.relname, c.relrowsecurity, c.relforcerowsecurity "
                "FROM pg_class c JOIN pg_namespace n ON n.oid = "
                "c.relnamespace WHERE n.nspname IN ('business', "
                "'qualification') AND c.relkind = 'r' AND c.relname IN ("
                "'organization_financial_score_observations', "
                "'organization_financial_score_state', 'crm_settings') "
                "ORDER BY c.relname")
            tables = [dict(r) for r in cur.fetchall()]
            cur.execute(
                "SELECT tablename, policyname, permissive, roles::text, "
                "cmd FROM pg_policies WHERE schemaname IN ('business', "
                "'qualification') AND tablename IN ("
                "'organization_financial_score_observations', "
                "'organization_financial_score_state', 'crm_settings') "
                "ORDER BY tablename")
            policies = [dict(r) for r in cur.fetchall()]
            cur.execute(
                "SELECT pg_get_constraintdef(oid) AS def FROM pg_constraint "
                "WHERE conname = "
                "'organization_identifiers_identifier_type_check'")
            ident_check = (cur.fetchone() or {}).get("def")
    return {"tables": tables, "policies": policies,
            "identifier_check": ident_check}


# ---------------------------------------------------------------------------
# SQLite intelligence-store fixtures (assessments + archived drafts)
# ---------------------------------------------------------------------------

def _gate_payload(outcome: str, *, score: Optional[str],
                  threshold: str = "0.6", provider: str = "CREDITSAFE",
                  policy_version: int = 1, observed_at: str =
                  "2026-10-01T00:00:00+00:00") -> Dict[str, Any]:
    gate = {
        "outcome": outcome,
        "score": score,
        "threshold": threshold,
        "provider": provider,
        "fresh": True,
        "policy_version": policy_version,
        "reason_code": {"pass": "financial_gate_passed",
                        "fail": "financial_gate_failed",
                        "unknown": "financial_gate_unknown"}[outcome],
        "observed_at": observed_at,
        "evaluated_at": "2026-10-10T09:00:00+00:00",
    }
    return gate


def cmd_seed_assessments(args) -> int:
    from retriva_crm_assistant.company_intelligence import (
        get_intelligence_store,
    )
    store = get_intelligence_store()
    suffix = args.suffix or uuid.uuid4().hex[:8]
    ids: Dict[str, Any] = {"suffix": suffix}
    specs = {
        "pass": {"candidate_id": f"cand_pass_{suffix}",
                 "candidate_name": "Synthetic Pass Srl",
                 "tier": "strong_fit",
                 "financial_gate": _gate_payload("pass", score="0.8")},
        "pass2": {"candidate_id": f"cand_pass2_{suffix}",
                  "candidate_name": "Synthetic Pass2 Srl",
                  "tier": "strong_fit",
                  "financial_gate": _gate_payload("pass", score="0.8")},
        "fail": {"candidate_id": f"cand_fail_{suffix}",
                 "candidate_name": "Synthetic Fail Srl",
                 "tier": "strong_fit",
                 "financial_gate": _gate_payload("fail", score="0.4")},
        "unknown": {"candidate_id": f"cand_unknown_{suffix}",
                    "candidate_name": "Synthetic Unknown Srl",
                    "tier": "possible_fit",
                    "financial_gate": _gate_payload("unknown",
                                                    score=None)},
        "legacy": {"candidate_id": f"cand_legacy_{suffix}",
                   "candidate_name": "Synthetic Legacy Srl",
                   "tier": "strong_fit"},
        "tampered": {"candidate_id": f"cand_tampered_{suffix}",
                     "candidate_name": "Synthetic Tampered Srl",
                     "tier": "strong_fit",
                     "financial_gate": _gate_payload("pass",
                                                     score="0.5")},
        "mismatch": {"candidate_id": f"cand_mismatch_{suffix}",
                     "candidate_name": "Synthetic Mismatch Srl",
                     "tier": "strong_fit",
                     "financial_gate": _gate_payload("pass",
                                                     score="0.8")},
    }
    for key, payload in specs.items():
        payload = dict(payload)
        payload["acp_id"] = f"acp_{suffix}"
        ident = store.upsert_identity(TENANT_A,
                                      payload["candidate_name"])
        rec = store.create_assessment(
            TENANT_A, ident.company_identity_id, created_by="pipeline",
            pipeline_payload=payload)
        ids[key] = {"assessment_id": rec.assessment_id,
                    "company_identity_id": ident.company_identity_id}

    # Legacy assessments (no financial-gate snapshot) bound to the
    # synthetic PG organizations A4 (current fresh score) and A3 (no
    # score) for re-evaluation acceptance.
    for key, org, name in (
            ("legacy_a4", A4, "Synthetic Delta Legacy Srl"),
            ("legacy_a3", A3, "Synthetic Gamma Legacy Srl")):
        ident = store.upsert_identity(TENANT_A, name)
        rec = store.create_assessment(
            TENANT_A, ident.company_identity_id, created_by="pipeline",
            pipeline_payload={"candidate_id": f"cand_{key}_{suffix}",
                              "candidate_name": name,
                              "acp_id": f"acp_{suffix}",
                              "tier": "possible_fit"})
        ids[key] = {"assessment_id": rec.assessment_id,
                    "company_identity_id": ident.company_identity_id,
                    "organization_id": org}

    # Archived job with drafts (pass + fail) for the archived-draft seam.
    from retriva_crm_assistant.jobs import (
        QualificationJob as Job, JobState,
    )
    from retriva_crm_assistant.models import (
        QualificationReport, QualificationResult,
    )
    from retriva_crm_assistant.job_archive import get_job_archive
    report = QualificationReport(
        session_id=f"ses_s037acc_{suffix}", kb_id="kb_synthetic",
        results=[
            QualificationResult.model_validate({
                "candidate_id": f"cand_arch_pass_{suffix}",
                "candidate_name": "Synthetic Archived Pass Srl",
                "acp_id": f"acp_arch_{suffix}", "tier": "strong_fit",
                "financial_gate": _gate_payload("pass", score="0.7")}),
            QualificationResult.model_validate({
                "candidate_id": f"cand_arch_fail_{suffix}",
                "candidate_name": "Synthetic Archived Fail Srl",
                "acp_id": f"acp_arch_{suffix}", "tier": "strong_fit",
                "financial_gate": _gate_payload("fail", score="0.3")}),
        ])
    session_id = f"ses_s037acc_{suffix}"
    job = Job(session_id=session_id, attachment_id=f"att_{suffix}",
              kb_id="kb_synthetic")
    job.state = JobState.COMPLETED
    job.completed_at = datetime.now(UTC).isoformat()
    job_id = job.job_id
    archive = get_job_archive()
    result = archive.archive_completed_job(
        job, report=report, tenant_id=TENANT_A)
    listing = archive.list_drafts(TENANT_A, job_id=job_id)
    draft_records = listing.get("drafts", []) if isinstance(
        listing, dict) else list(listing)

    def _draft_fields(d: Any) -> Dict[str, Any]:
        get = (d.get if isinstance(d, dict)
               else lambda k, _d=d: getattr(_d, k, None))
        return {"draft_id": get("draft_id"),
                "candidate_result_id": get("candidate_result_id"),
                "candidate_name": get("candidate_name"),
                "review_status": get("review_status")}

    ids["archive"] = {
        "job_id": job_id,
        "archive_result": result,
        "drafts": [_draft_fields(d) for d in draft_records],
    }
    out(ids)
    return 0


def cmd_refresh_a4(_args) -> int:
    from retriva_crm_assistant.postgres.qualification.enrichment import (
        creditsafe_stage as stage,
    )
    from retriva_crm_assistant.postgres.qualification import (
        financial_scores as fstore,
    )
    p, _ = pool()
    mock = success_mock(connect_id="CS-S037-A4", reg_no="RN-S037-A4",
                        name="Synthetic Delta Srl", value=60)
    provider, guard = make_mock_provider(mock)
    r = stage.run_creditsafe_stage(p, TENANT_A, A4, provider=provider,
                                   forced_refresh=True)
    out({"scenario": "refresh-a4-to-0.6", "outcome": r.to_result(),
         "mock_calls": mock.summary(),
         "budget_snapshot": guard.snapshot(),
         "state": state_to_dict(fstore.get_state(p, TENANT_A, A4)),
         "current": obs_to_dict(
             fstore.get_current_observation(p, TENANT_A, A4)),
         "history_len": len(fstore.list_observations(p, TENANT_A, A4))})
    return 0


def cmd_report_artifacts(args) -> int:
    import io as _io
    import os as _os
    from openpyxl import load_workbook

    from retriva_crm_assistant.models import (
        FinancialGateOutcome, FinancialGateResult, FinancialScoreStatus,
        QualificationReport, QualificationResult, QualificationTier,
    )
    from retriva_crm_assistant.report import (
        build_markdown_report, build_xlsx_report,
    )

    def _gate(outcome: str, score: Optional[str]) -> FinancialGateResult:
        return FinancialGateResult(
            outcome=FinancialGateOutcome(outcome),
            score=Decimal(score) if score is not None else None,
            threshold=Decimal("0.6"), provider="CREDITSAFE",
            observation_id="fsobs_" + outcome, fresh=True,
            evaluated_at="2026-10-10T13:00:00+00:00",
            reason_code={
                "pass": "financial_gate_passed",
                "fail": "financial_gate_failed",
                "unknown": "financial_gate_unknown",
            }[outcome],
            score_status=(FinancialScoreStatus.REPORT_UNAVAILABLE
                          if outcome == "unknown"
                          else FinancialScoreStatus.AVAILABLE),
            policy_version=1)

    results = [
        QualificationResult(
            candidate_id="cand_art_pass",
            candidate_name="Synthetic Artifact Pass Srl", acp_id="acp_p5",
            tier=QualificationTier.STRONG_FIT,
            financial_gate=_gate("pass", "0.8")),
        QualificationResult(
            candidate_id="cand_art_fail",
            candidate_name="Synthetic Artifact Fail Srl", acp_id="acp_p5",
            tier=QualificationTier.NOT_A_FIT,
            financial_gate=_gate("fail", "0.4")),
        QualificationResult(
            candidate_id="cand_art_unknown",
            candidate_name="Synthetic Artifact Unknown Srl", acp_id="acp_p5",
            tier=QualificationTier.INSUFFICIENT_EVIDENCE,
            financial_gate=_gate("unknown", None)),
        QualificationResult(
            candidate_id="cand_art_legacy",
            candidate_name="Synthetic Artifact Legacy Srl", acp_id="acp_p5",
            tier=QualificationTier.POSSIBLE_FIT),
    ]
    report = QualificationReport(
        session_id="ses_s037acc_artifacts", kb_id="kb_synthetic",
        results=results)
    md = build_markdown_report(report)
    xlsx = build_xlsx_report(report, include_diagnostics=False)

    outdir = getattr(args, "outdir", None) or "/tmp/sp037acc-artifacts"
    _os.makedirs(outdir, exist_ok=True)
    md_path = _os.path.join(outdir, "qualification_report.md")
    xlsx_path = _os.path.join(outdir, "qualification_report.xlsx")
    with open(md_path, "w") as fh:
        fh.write(md)
    with open(xlsx_path, "wb") as fh:
        fh.write(xlsx)

    wb = load_workbook(_io.BytesIO(xlsx))
    rows = [[str(c) if c is not None else "" for c in row]
            for row in wb["Financial qualification"].values]

    leak_patterns = ["CS-S037", "CS-V009", "Bearer", "Authorization",
                     "synthetic-user", "synthetic-pass", "synthetic.jwt"]
    blob = (md + "\n" + xlsx.decode("latin-1", "ignore"))
    leaks = sorted({p for p in leak_patterns if p in blob})

    out({
        "md_path": md_path,
        "xlsx_path": xlsx_path,
        "sheetnames": wb.sheetnames,
        "financial_sheet_rows": rows,
        "md_sections": [l for l in md.splitlines()
                        if l.startswith("### Financial qualification")],
        "leak_patterns_found": leaks,
        "checks": {
            "md_has_per_candidate_section": (
                md.count("### Financial qualification") == len(results)),
            "md_has_reason_codes": all(
                f"- Reason code: financial_gate_{o}" in md
                for o in ("passed", "failed", "unknown")),
            "md_has_eval_timestamp": md.count(
                "- Evaluation timestamp: 2026-10-10T13:00:00+00:00")
            == 3,
            "md_has_policy_version": md.count("- Policy version: 1") == 3,
            "md_marks_legacy": (
                md.count("Legacy artifact") == 1),
            "xlsx_sheet_present": (
                "Financial qualification" in wb.sheetnames),
            "xlsx_rows": len(rows) - 3,  # header block rows
            "no_leaks_in_artifacts": not leaks,
        },
    })
    return 0


def cmd_seed_cohort(_args) -> int:
    import dataclasses

    def _member_dict(m: Any) -> Dict[str, Any]:
        if hasattr(m, "model_dump"):
            return m.model_dump()
        if dataclasses.is_dataclass(m):
            return dataclasses.asdict(m)
        return dict(getattr(m, "__dict__", {}))

    from retriva_crm_assistant.postgres.qualification import (
        service as qual_service,
    )
    from retriva_crm_assistant.postgres.qualification import (
        repository as repo,
    )
    p, _ = pool()
    result = qual_service.propose_cohort(
        p, TENANT_A, actor_id="crm-admin",
        name="Synthetic Spec 037 Acceptance Cohort",
        description="Synthetic cohort for Phase 5 deployed acceptance")
    version = result.get("version")
    version_id = getattr(version, "cohort_version_id", None)
    cohort_id = getattr(version, "cohort_id", None)
    members = []
    if version_id:
        members = [_member_dict(m) for m in repo.list_cohort_members(
            p, TENANT_A, version_id)]
    out({"cohort_id": cohort_id,
         "cohort_version_id": version_id,
         "counts": result.get("counts"),
         "proposed_now": result.get("proposed_now"),
         "member_count": len(members),
         "members": members})
    return 0


def main(argv: Optional[List[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        prog="sp037acc_harness", description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("seed-fixtures")
    sub.add_parser("scenario-success")
    sub.add_parser("scenario-zero-budget")
    sub.add_parser("scenario-budget-boundary")
    sub.add_parser("scenario-failure-recover")
    sub.add_parser("scenario-conflicts")
    sub.add_parser("scenario-stale-and-tenant-isolation")
    q = sub.add_parser("query")
    q.add_argument("what", choices=["threshold", "states", "rls", "audit", "provider"])
    q.add_argument("--tenant", default=None)
    a = sub.add_parser("seed-assessments")
    a.add_argument("--suffix", default=None)
    sub.add_parser("seed-cohort")
    sub.add_parser("scenario-refresh-a4")
    ra = sub.add_parser("report-artifacts")
    ra.add_argument("--outdir", default=None)
    args = parser.parse_args(argv)
    handlers = {
        "seed-fixtures": cmd_seed_fixtures,
        "scenario-success": cmd_scenario_success,
        "scenario-zero-budget": cmd_scenario_zero_budget,
        "scenario-budget-boundary": cmd_scenario_budget_boundary,
        "scenario-failure-recover": cmd_scenario_failure_recover,
        "scenario-conflicts": cmd_scenario_conflicts,
        "scenario-stale-and-tenant-isolation":
            cmd_scenario_stale_and_tenant_isolation,
        "query": cmd_query,
        "seed-assessments": cmd_seed_assessments,
        "seed-cohort": cmd_seed_cohort,
        "report-artifacts": cmd_report_artifacts,
        "scenario-refresh-a4": cmd_refresh_a4,
    }
    return handlers[args.command](args)


if __name__ == "__main__":
    raise SystemExit(main())
